from docx import Document


class DocumentGenerator:

    def generate(self, proposal, output_path):

        doc = Document()

        doc.add_heading(
            'Project Proposal',
            level=1
        )

        doc.add_heading(
            'Executive Summary',
            level=2
        )
        doc.add_paragraph(
            proposal.executive_summary
        )

        doc.add_heading(
            'Proposed Solution',
            level=2
        )
        doc.add_paragraph(
            proposal.proposed_solution
        )

        doc.add_heading(
            'Technology Stack',
            level=2
        )

        for tech in proposal.technology_stack:
            doc.add_paragraph(
                tech,
                style='List Bullet'
            )

        doc.add_heading(
            'Milestones',
            level=2
        )

        for milestone in proposal.milestones:
            doc.add_paragraph(
                milestone,
                style='List Bullet'
            )

        doc.add_heading(
            'Estimated Budget',
            level=2
        )
        doc.add_paragraph(
            proposal.estimated_budget
        )

        doc.add_heading(
            'Risks',
            level=2
        )

        for risk in proposal.risks:
            doc.add_paragraph(
                risk,
                style='List Bullet'
            )

        doc.save(output_path)